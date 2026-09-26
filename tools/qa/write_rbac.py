"""Write the RBAC run (out/rbac_results.json) into the RBAC Tests sheet of the
extended coverage workbook: Status, Actual Result, Defect ID, Tester and the
evidence screenshot. Nothing else in the sheet changes. The Defect Log is
written by write_results.py, from defects.py.

    python3 write_rbac.py
"""
import json, os
import openpyxl
from defects import DEFECTS

HERE = os.path.dirname(os.path.abspath(__file__))
EXT = os.path.join(os.path.dirname(os.path.dirname(HERE)), "School_ERP_Functional_Testing_Extended_Coverage.xlsx")
results = {r["id"]: r for r in json.load(open(os.path.join(HERE, "out", "rbac_results.json")))}
defect_of = {s: d for d in DEFECTS for s in d.get("rbac", [])}

wb = openpyxl.load_workbook(EXT)
ws = wb["RBAC Tests"]
col = {c.value: c.column for c in ws[1]}
for row in range(2, ws.max_row + 1):
    r = results.get(ws.cell(row, 1).value)
    if not r:
        continue
    d = defect_of.get(r["scr"]) if r["scenario"] == "authorized" else None
    actual = r["actual"]
    if d and d["status"] == "Fixed":
        how = ('blocked', 'test plan corrected') if d['title'].startswith('Test plan') else ('failed', 'fixed')
        actual = f"First run {how[0]} ({d['id']}: {d['title'][0].lower() + d['title'][1:]}); {how[1]} and run again: {actual}"
    ws.cell(row, col["Status"]).value = r["status"]
    ws.cell(row, col["Actual Result"]).value = actual
    ws.cell(row, col["Defect ID"]).value = d["id"] if d else None
    ws.cell(row, col["Tester"]).value = r["tester"]
    ws.cell(row, col["Evidence / Comments"]).value = f"Screenshot: {r['evidence']}"
wb.save(EXT)
st = [r["status"] for r in results.values()]
print({s: st.count(s) for s in sorted(set(st))}, "written to RBAC Tests")
