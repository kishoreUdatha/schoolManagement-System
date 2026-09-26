"""Correct the persona of the screens the plan assigned to a role that does not
use them (DEF-004, DEF-005, DEF-009), in every sheet of both workbooks that
names a role: the role column, and "Sign in as <role>" in the steps. The API
sheet's Authorization column lists every role allowed and is left alone.

    python3 fix_personas.py
"""
import os, re
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
BOOKS = ["School_ERP_Functional_Test_Plan_and_Cases.xlsx", "School_ERP_Functional_Testing_Extended_Coverage.xlsx"]
ROLE_COLUMNS = {"Role / Persona", "Role", "Business Role", "Test User / Role"}
SKIP_SHEETS = {"API Functional Tests"}
PERSONA = {
    # the school's own setup: profile, branches, academic years, terms, setup
    **{f"SCR-0{n}": "School Admin" for n in range(24, 33)},
    "SCR-114": "Parent",              # the parent asks leave for their child
    "SCR-116": "School Admin",        # the office reviews late arrivals and early exits
    "SCR-146": "Exam Coordinator",    # practical/internal marks: the Examinations job
    "SCR-148": "Exam Coordinator",    # marks sign-off: the Examinations job
    "SCR-223": "Discipline In-charge",
    "SCR-224": "Discipline In-charge",
    "SCR-258": "School Admin",        # the office keeps the document register
}

changed = {}
for name in BOOKS:
    path = os.path.join(ROOT, name)
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        if ws.title in SKIP_SHEETS or ws.max_row < 2:
            continue
        hdr = {c.value: c.column for c in ws[1] if c.value}
        scr_col = hdr.get("Screen ID")
        role_col = next((hdr[h] for h in ROLE_COLUMNS if h in hdr), None)
        if not scr_col or not role_col:
            continue
        for row in range(2, ws.max_row + 1):
            new = PERSONA.get(ws.cell(row, scr_col).value)
            old = ws.cell(row, role_col).value
            # (an RBAC "unauthorized" row names no role: "User without required role…")
            if not new or not old or old == new or old.startswith("User without"):
                continue
            ws.cell(row, role_col).value = new
            pat = re.compile(rf"\bas {re.escape(old)}\b")
            for col in range(1, ws.max_column + 1):
                c = ws.cell(row, col)
                if col != role_col and isinstance(c.value, str) and pat.search(c.value):
                    c.value = pat.sub(f"as {new}", c.value)
            changed[(name[:40], ws.title)] = changed.get((name[:40], ws.title), 0) + 1
    wb.save(path)
for (book, sheet), n in changed.items():
    print(f"{book} / {sheet}: {n} rows")
